"""File intake: storage, filtering (FR-5.5), text acquisition (spec 5.7 step 1)."""

import hashlib
import ipaddress
import mimetypes
import os
import re
import socket
from pathlib import Path
from urllib.parse import urlparse

import httpx

from app.config import settings

ALLOWED_EXTENSIONS = {
    "pdf", "png", "jpg", "jpeg", "tif", "tiff",
    "docx", "xlsx", "xls", "csv", "txt",
}
TINY_IMAGE_BYTES = 10 * 1024
MIN_PDF_CHARS_PER_PAGE = 50
MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024


def file_store_root() -> Path:
    root = Path(settings.file_store_dir)
    root.mkdir(parents=True, exist_ok=True)
    return root


def sanitize_filename(name: str) -> str:
    base = os.path.basename(name or "file")
    base = re.sub(r"[^A-Za-z0-9._-]", "_", base).strip("._") or "file"
    return base[:200]


def store_upload(filename: str, content: bytes, submission_id: str) -> dict:
    """Store bytes under a generated name. Returns file record fields."""
    safe = sanitize_filename(filename)
    ext = safe.rsplit(".", 1)[-1].lower() if "." in safe else ""
    import uuid

    stored = f"{uuid.uuid4().hex}.{ext}" if ext else uuid.uuid4().hex
    directory = file_store_root() / submission_id
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / stored
    path.write_bytes(content)
    mime, _ = mimetypes.guess_type(safe)
    return {
        "filename": safe,
        "mime": mime or "application/octet-stream",
        "size": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
        "storage_path": str(path),
    }


def filter_file(filename: str, size: int, max_mb: int | None = None) -> tuple[bool, str]:
    """Pre-AI filtering (FR-5.5). Returns (accepted, skip_reason)."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        return False, f"unsupported_type: .{ext or '?'}"
    limit = (max_mb if max_mb is not None else settings.max_extract_file_mb) * 1024 * 1024
    if size > limit:
        return False, f"oversize: {size} bytes"
    if ext in ("png", "jpg", "jpeg", "tif", "tiff") and size < TINY_IMAGE_BYTES:
        return False, "tiny_image: likely signature/logo"
    return True, ""


def acquire_text(storage_path: str, mime: str) -> tuple[str, str, str]:
    """Extract text; returns (text, text_source, text_path).

    text_source is one of: pdf_text, scanned, docx, xlsx, csv, txt, unsupported.
    Scanned PDFs/images return empty text — the extraction service routes them
    through the vision OCR provider or fails them for manual entry.
    """
    path = Path(storage_path)
    if not path.exists():
        return "", "unsupported", ""

    text, source = "", "unsupported"
    try:
        if mime == "application/pdf" or path.suffix.lower() == ".pdf":
            text, source = _pdf_text(path)
        elif path.suffix.lower() == ".docx":
            text, source = _docx_text(path), "docx"
        elif path.suffix.lower() in (".xlsx", ".xls"):
            text, source = _xlsx_text(path), "xlsx"
        elif path.suffix.lower() == ".csv" or mime == "text/csv":
            text, source = path.read_text(errors="replace"), "csv"
        elif path.suffix.lower() == ".txt" or (mime or "").startswith("text/"):
            text, source = path.read_text(errors="replace"), "txt"
        elif (mime or "").startswith("image/"):
            source = "scanned"
    except Exception:
        return "", "unsupported", ""

    text_path = ""
    if text:
        text_path = str(path) + ".txt"
        Path(text_path).write_text(text, errors="replace")
    return text, source, text_path


def _pdf_text(path: Path) -> tuple[str, str]:
    import pymupdf

    doc = pymupdf.open(path)
    try:
        pages = [page.get_text().strip() for page in doc]
    finally:
        doc.close()
    non_empty = [(i + 1, text) for i, text in enumerate(pages) if text]
    if not non_empty:
        return "", "scanned"
    if len(non_empty) == 1:
        joined = non_empty[0][1]
    else:
        # Multi-page: mark page boundaries so the model can follow tables
        # continued across pages and ignore repeated headers/footers.
        total = len(pages)
        joined = "\n\n".join(
            f"=== PAGE {num} of {total} ===\n{text}" for num, text in non_empty
        )
    joined = joined.strip()
    per_page = len(joined) / max(len(pages), 1)
    if per_page < MIN_PDF_CHARS_PER_PAGE:
        return "", "scanned"
    return joined, "pdf_text"


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(path)
    return "\n".join(p.text for p in document.paragraphs).strip()


def _xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines).strip()


def estimate_tokens(text: str) -> int:
    return max(1, len(text) // 4)


def is_private_url(url: str) -> bool:
    """True if the URL targets a non-public address (FR-S3)."""
    try:
        parsed = urlparse(url)
    except Exception:
        return True
    if parsed.scheme not in ("http", "https"):
        return True
    if parsed.username or parsed.password:
        return True
    host = (parsed.hostname or "").lower()
    if host in ("localhost",):
        return True
    try:
        infos = socket.getaddrinfo(host, None)
    except OSError:
        return True
    for info in infos:
        try:
            ip = ipaddress.ip_address(info[4][0])
        except ValueError:
            return True
        if not ip.is_global:
            return True
    return False


def validate_provider_url(base_url: str, is_local: bool) -> None:
    """Reject private base URLs unless the provider is explicitly local (FR-S3)."""
    if not base_url or urlparse(base_url).scheme not in ("http", "https"):
        raise ValueError("base_url must be an http(s) URL")
    if is_private_url(base_url) and not is_local:
        raise ValueError("base URL targets a private address; mark the provider local or use a public URL")


async def fetch_url(url: str, max_bytes: int = MAX_DOWNLOAD_BYTES) -> tuple[str, bytes]:
    """Download a document URL with SSRF guard, timeout, and size cap."""
    if is_private_url(url):
        raise ValueError("URL targets a private or invalid address")
    try:
        async with httpx.AsyncClient(timeout=30, follow_redirects=True, max_redirects=3) as client:
            async with client.stream("GET", url) as resp:
                if resp.status_code != 200:
                    raise ValueError(f"Download failed with status {resp.status_code}")
                chunks: list[bytes] = []
                total = 0
                async for chunk in resp.aiter_bytes(65536):
                    total += len(chunk)
                    if total > max_bytes:
                        raise ValueError("Download exceeds size cap")
                    chunks.append(chunk)
                content = b"".join(chunks)
    except httpx.HTTPError as exc:
        raise ValueError(f"Download failed: {exc}") from exc
    parsed = urlparse(url)
    filename = os.path.basename(parsed.path.rstrip("/")) or "download"
    ctype = resp.headers.get("content-type", "").split(";")[0].strip()
    if "." not in filename and ctype:
        ext = mimetypes.guess_extension(ctype) or ""
        filename += ext
    return filename, content
